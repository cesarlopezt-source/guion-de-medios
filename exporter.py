import html
import re
from datetime import datetime
from io import BytesIO

import openpyxl
from docx import Document
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from config.paises import PAISES

PAIS_HEX = {
    "SV": "D6E4F5", "GT": "D9F0E3", "HN": "FCE8D5", "NI": "EADCF2",
    "CR": "FFF3C4", "PA": "FAD9D9", "US": "DCDCF0", "REG": "E5E8EB",
}
_MARCA = re.compile(r"^\s*\[([^\]]+)\]\s*$")
_ENLACE = re.compile(r"^-\s*(.+?):\s*(https?://\S+)\s*$")  # "- Medio: https://..." en el mini brief
NOMBRE = "MIXCOM News Monitor"


def _estrellas(n: int) -> str:
    return "★" * n + "☆" * (5 - n)


def _fecha() -> str:
    return datetime.now().strftime("%d/%m/%Y %H:%M")


def _por_pais(items: list) -> list:
    """[(nombre_pais, [items])] en el orden de PAISES."""
    grupos = {}
    for item in items:
        grupos.setdefault(item["pais"], []).append(item)
    return [(PAISES[p], grupos[p]) for p in PAISES if p in grupos]


def _link(url: str, texto: str = None) -> str:
    """Enlace clickeable para un Paragraph de reportlab."""
    return f'<a href="{html.escape(url, quote=True)}" color="blue">{html.escape(texto or url)}</a>'


def _bytes(guardar) -> bytes:
    buf = BytesIO()
    guardar(buf)
    return buf.getvalue()


def export_word(items: list) -> bytes:
    doc = Document()
    doc.add_heading(f"{NOMBRE} · Informe de noticias", level=0)
    doc.add_paragraph(f"Generado: {_fecha()}")
    for pais, noticias in _por_pais(items):
        doc.add_heading(pais, level=1)
        for item in noticias:
            doc.add_heading(f"{_estrellas(item['estrellas'])}  {item['titulo']}", level=2)
            doc.add_paragraph(f"Temas: {', '.join(item['temas'])}")
            doc.add_paragraph(item.get("resumen") or "Sin resumen.")
            if (item.get("valoracion") or "").strip():
                doc.add_paragraph(f"Valoración: {item['valoracion'].strip()}")
            doc.add_paragraph(f"Medios: {item['n_fuentes']}  |  Publicado: {item['publicado'][:16]}")
            for f in item["fuentes"]:
                doc.add_paragraph(f"{f['nombre']}: {f['url']}", style="List Bullet")
    return _bytes(doc.save)


def export_excel(items: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Noticias"
    ws.append(["País", "Temas", "Estrellas", "Título", "Resumen", "Valoración", "Medios", "Publicado", "Fuentes"])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2C3E50")
    for _, noticias in _por_pais(items):
        for item in noticias:
            ws.append([
                PAISES[item["pais"]], ", ".join(item["temas"]), _estrellas(item["estrellas"]), item["titulo"],
                item.get("resumen", ""), item.get("valoracion", ""), item["n_fuentes"], item["publicado"][:16],
                " | ".join(f"{f['nombre']}: {f['url']}" for f in item["fuentes"]),
            ])
            relleno = PatternFill("solid", fgColor=PAIS_HEX.get(item["pais"], "FFFFFF"))
            for celda in ws[ws.max_row]:
                celda.fill = relleno
    ws.auto_filter.ref = ws.dimensions
    for col, ancho in {"B": 25, "D": 55, "E": 65, "F": 55, "I": 80}.items():
        ws.column_dimensions[col].width = ancho
    return _bytes(wb.save)


def _pdf(construir) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    doc.build(construir(getSampleStyleSheet()))
    return buf.getvalue()


def export_pdf(items: list) -> bytes:
    def construir(st):
        e = html.escape
        story = [Paragraph(f"{NOMBRE} · Informe de noticias", st["Title"]),
                 Paragraph(f"Generado: {_fecha()}", st["Normal"]), Spacer(1, 0.5 * cm)]
        for pais, noticias in _por_pais(items):
            story.append(Paragraph(e(pais), st["Heading1"]))
            for item in noticias:
                story.append(Paragraph(e(f"{_estrellas(item['estrellas'])}  {item['titulo']}"), st["Heading2"]))
                story.append(Paragraph(f"<i>{e(', '.join(item['temas']))}</i>", st["Normal"]))
                story.append(Paragraph(e(item.get("resumen") or "Sin resumen."), st["Normal"]))
                if (item.get("valoracion") or "").strip():
                    story.append(Paragraph(f"<b>Valoración:</b> {e(item['valoracion'].strip())}", st["Normal"]))
                story.append(Paragraph(f"<b>Medios:</b> {item['n_fuentes']}  |  {e(item['publicado'][:16])}",
                                       st["Normal"]))
                for f in item["fuentes"]:
                    story.append(Paragraph(f"• {e(f['nombre'])}: {_link(f['url'])}", st["Normal"]))
                story.append(Spacer(1, 0.4 * cm))
        return story
    return _pdf(construir)


def _titulo_guion() -> str:
    return f"Guion de medios · {datetime.now().strftime('%d/%m/%Y')}"


def export_guion_word(texto: str) -> bytes:
    doc = Document()
    doc.add_heading(_titulo_guion(), level=0)
    for linea in texto.splitlines():
        marca = _MARCA.match(linea)
        if marca:
            doc.add_heading(marca.group(1), level=2)
        elif linea.strip():
            doc.add_paragraph(linea.strip())
    return _bytes(doc.save)


def export_guion_pdf(texto: str) -> bytes:
    def construir(st):
        story = [Paragraph(html.escape(_titulo_guion()), st["Title"])]
        for linea in texto.splitlines():
            marca = _MARCA.match(linea)
            if marca:
                story.append(Paragraph(html.escape(marca.group(1)), st["Heading2"]))
            elif linea.strip():
                story.append(Paragraph(html.escape(linea.strip()), st["Normal"]))
                story.append(Spacer(1, 0.2 * cm))
        return story
    return _pdf(construir)


# ---------- Solo temas y links ----------

def export_links_word(items: list) -> bytes:
    doc = Document()
    doc.add_heading(f"{NOMBRE} · Temas y links", level=0)
    doc.add_paragraph(f"Generado: {_fecha()}")
    for item in items:
        doc.add_heading(item["titulo"], level=2)
        for f in item["fuentes"]:
            doc.add_paragraph(f"{f['nombre']}: {f['url']}", style="List Bullet")
    return _bytes(doc.save)


def export_links_excel(items: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Temas y links"
    ws.append(["Tema", "Medio", "Link"])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2C3E50")
    for item in items:
        for i, f in enumerate(item["fuentes"]):
            ws.append([item["titulo"] if i == 0 else "", f["nombre"], f["url"]])
            ws.cell(ws.max_row, 3).hyperlink = f["url"]
    for col, ancho in {"A": 70, "B": 25, "C": 90}.items():
        ws.column_dimensions[col].width = ancho
    return _bytes(wb.save)


def export_links_pdf(items: list) -> bytes:
    def construir(st):
        e = html.escape
        story = [Paragraph(f"{NOMBRE} · Temas y links", st["Title"]),
                 Paragraph(f"Generado: {_fecha()}", st["Normal"]), Spacer(1, 0.5 * cm)]
        for item in items:
            story.append(Paragraph(e(item["titulo"]), st["Heading3"]))
            for f in item["fuentes"]:
                story.append(Paragraph(f"• {e(f['nombre'])}: {_link(f['url'])}", st["Normal"]))
            story.append(Spacer(1, 0.3 * cm))
        return story
    return _pdf(construir)


# ---------- Mini brief ----------

def _titulo_brief() -> str:
    return f"{NOMBRE} · Mini brief · {datetime.now().strftime('%d/%m/%Y')}"


def export_brief_word(texto: str) -> bytes:
    doc = Document()
    doc.add_heading(_titulo_brief(), level=0)
    for linea in texto.splitlines():
        marca, enlace = _MARCA.match(linea), _ENLACE.match(linea.strip())
        if marca:
            doc.add_heading(marca.group(1), level=2)
        elif enlace:
            doc.add_paragraph(f"{enlace.group(1)}: {enlace.group(2)}", style="List Bullet")
        elif linea.strip():
            doc.add_paragraph(linea.strip())
    return _bytes(doc.save)


def export_brief_pdf(texto: str) -> bytes:
    def construir(st):
        e = html.escape
        story = [Paragraph(e(_titulo_brief()), st["Title"])]
        anterior_marca = False
        for linea in texto.splitlines():
            marca, enlace = _MARCA.match(linea), _ENLACE.match(linea.strip())
            if marca:
                story.append(Paragraph(e(marca.group(1)), st["Heading2"]))
            elif enlace:
                story.append(Paragraph(f"• {e(enlace.group(1))}: {_link(enlace.group(2))}", st["Normal"]))
            elif linea.strip():
                # La línea que sigue a la marca de una noticia es su titular.
                estilo = st["Heading4"] if anterior_marca else st["Normal"]
                story.append(Paragraph(e(linea.strip()), estilo))
                story.append(Spacer(1, 0.15 * cm))
            anterior_marca = bool(marca)
        return story
    return _pdf(construir)
