from io import BytesIO

import openpyxl

import exporter
from tests.fakes import GUION_OK, N1, N2


def test_export_excel_columnas_y_orden_por_pais():
    wb = openpyxl.load_workbook(BytesIO(exporter.export_excel([N2, {**N1, "valoracion": "V"}])))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["País", "Temas", "Estrellas", "Título", "Resumen", "Valoración",
                                        "Medios", "Publicado", "Fuentes"]
    assert ws["A2"].value == "El Salvador" and ws["F2"].value == "V" and ws["B3"].value == "Economía, Inversión"


def test_export_word_y_pdf_noticias():
    assert exporter.export_word([N1, N2])[:2] == b"PK"
    assert exporter.export_pdf([N1, N2])[:4] == b"%PDF"


def test_export_guion_word_y_pdf():
    assert exporter.export_guion_word(GUION_OK)[:2] == b"PK"
    assert exporter.export_guion_pdf(GUION_OK)[:4] == b"%PDF"


def test_export_guion_word_marcas_como_subtitulos():
    from docx import Document
    doc = Document(BytesIO(exporter.export_guion_word(GUION_OK)))
    titulos = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "APERTURA" in titulos and "NOTICIA 1 · El Salvador" in titulos


def test_export_temas_y_links():
    from docx import Document
    wb = openpyxl.load_workbook(BytesIO(exporter.export_links_excel([N1, N2])))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["Tema", "Medio", "Link"]
    assert ws["A2"].value == "Remesas crecen 5%" and ws["C2"].hyperlink.target == "https://lpg/1"
    doc = Document(BytesIO(exporter.export_links_word([N1])))
    assert "La Prensa Gráfica: https://lpg/1" in [p.text for p in doc.paragraphs]
    assert exporter.export_links_pdf([N1, N2])[:4] == b"%PDF"


def test_export_brief():
    from docx import Document
    texto = "[ANÁLISIS GENERAL]\nPanorama.\n\n[1 · El Salvador]\nRemesas\nBrief.\n- LPG: https://lpg/1"
    doc = Document(BytesIO(exporter.export_brief_word(texto)))
    titulos = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "ANÁLISIS GENERAL" in titulos and "1 · El Salvador" in titulos
    assert exporter.export_brief_pdf(texto)[:4] == b"%PDF"


def test_pdf_con_links_clickeables():
    pdf = exporter.export_pdf([{**N1, "fuentes": [{"nombre": "A&B", "url": "https://x.com/?a=1&b=2"}]}])
    assert b"/URI" in pdf
